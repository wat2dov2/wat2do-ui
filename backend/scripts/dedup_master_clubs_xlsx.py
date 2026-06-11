#!/usr/bin/env python3
"""Dedup duplicate (School, Name) rows in all_schools_student_clubs_master.xlsx.

The directory scrape occasionally surfaces the same club twice (typically
with a `/copy-of-<slug>` Directory URL).  This script collapses each
duplicate pair to a single row, preferring the variant whose Directory URL
does not contain ``copy-of`` so we keep the canonical listing.

Usage (from backend/):
  python scripts/dedup_master_clubs_xlsx.py            # dry-run report
  python scripts/dedup_master_clubs_xlsx.py --apply    # rewrite the workbook
"""

from __future__ import annotations

import argparse
import logging
from collections import defaultdict
from pathlib import Path

import openpyxl

log = logging.getLogger(__name__)

XLSX_PATH = (
    Path(__file__).resolve().parent.parent
    / "services"
    / "scraper"
    / "all_schools_student_clubs_master.xlsx"
)

_SCHOOL_COL = 1
_NAME_COL = 2
_DIRECTORY_URL_COL = 5


def _is_canonical_url(url: str | None) -> bool:
    """Return True if the directory URL does not look like a copy/duplicate."""
    if not url:
        return False
    return "copy-of" not in url.lower()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true",
                        help="Rewrite the workbook with duplicates removed.")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    # Group rows by (school, name).  Keep 1-based row indices because
    # openpyxl's ws.delete_rows uses them.
    groups: dict[tuple[str, str], list[int]] = defaultdict(list)
    for row_idx in range(3, ws.max_row + 1):
        school = ws.cell(row=row_idx, column=_SCHOOL_COL).value
        name = ws.cell(row=row_idx, column=_NAME_COL).value
        if not school or not name:
            continue
        groups[(str(school).strip(), str(name).strip())].append(row_idx)

    rows_to_delete: list[int] = []
    for key, indices in groups.items():
        if len(indices) <= 1:
            continue

        # Prefer the row whose Directory URL is canonical (no "copy-of").
        canonical = [
            i for i in indices
            if _is_canonical_url(ws.cell(row=i, column=_DIRECTORY_URL_COL).value)
        ]
        keep = canonical[0] if canonical else indices[0]
        drop = [i for i in indices if i != keep]
        rows_to_delete.extend(drop)
        school, name = key
        log.info("dup %r | %r — keep row %s, drop %s",
                 school, name, keep, drop)

    log.info("total duplicate groups: %s", sum(1 for v in groups.values() if len(v) > 1))
    log.info("total rows to delete: %s", len(rows_to_delete))

    if not args.apply:
        log.info("Dry run only.  Pass --apply to rewrite the workbook.")
        return 0

    # Delete in DESCENDING row order so earlier indices stay valid.
    for row_idx in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_idx, 1)

    wb.save(XLSX_PATH)
    log.info("Wrote %s (%s rows remaining)", XLSX_PATH, ws.max_row - 2)
    return 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    raise SystemExit(main())
