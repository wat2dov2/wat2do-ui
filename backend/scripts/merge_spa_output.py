#!/usr/bin/env python3
"""Merge ``spa_scrape.py`` JSON output into wat2do-clubs.xlsx.

Matches existing rows on normalized Directory URL, updates social links from
``--with-details`` scrapes, and appends clubs that are not already present.

Usage (from repo root):
  python backend/scripts/merge_spa_output.py
  python backend/scripts/merge_spa_output.py --input-dir /tmp/claude --schools ubc ucalgary ualberta
"""

from __future__ import annotations

import argparse
import json
import logging
import re
from pathlib import Path

import openpyxl

log = logging.getLogger(__name__)

XLSX_PATH = Path(__file__).resolve().parent.parent / "services" / "scraper" / "wat2do-clubs.xlsx"
DEFAULT_INPUT_DIR = Path("/tmp/claude")

COLUMNS = [
    "School",
    "Name",
    "Category",
    "Campus",
    "Directory URL",
    "Instagram URL",
    "Instagram Handle",
    "IG Source",
    "Discord URL",
    "Discord Source",
]

SPA_FILES = {
    "concordia": "spa-concordia.json",
    "memorial": "spa-memorial.json",
    "sfu": "spa-sfu.json",
    "ubc": "spa-ubc.json",
    "ucalgary": "spa-ucalgary.json",
    "ualberta": "spa-ualberta.json",
    "umanitoba": "spa-umanitoba.json",
}

# Keep curated directory research when it already beat a profile-page scrape.
_PROTECTED_IG_SOURCES = frozenset(
    {"found", "confirmed", "listing", "gap_found", "description_embedded"}
)


def _normalize_url(url: str | None) -> str:
    if not url:
        return ""
    return re.sub(r"/+$", "", str(url).strip().lower())


def _normalize_header(header: list[object | None]) -> list[str]:
    return ["" if value is None else str(value).strip() for value in header]


def _row_to_dict(values: tuple[object, ...]) -> dict[str, str | None]:
    padded = list(values) + [None] * len(COLUMNS)
    row = {col: padded[i] for i, col in enumerate(COLUMNS)}
    for key, value in row.items():
        if value is None:
            row[key] = None
        else:
            text = str(value).strip()
            row[key] = text or None
    return row


def _dict_to_row(row: dict[str, str | None]) -> list[str | None]:
    return [row.get(col) for col in COLUMNS]


def _should_update_ig(
    existing_source: str | None, spa_source: str | None, spa_ig: str | None
) -> bool:
    if not spa_ig or spa_source != "profile_page":
        return False
    base = (existing_source or "").split("|", 1)[0].strip().lower()
    if not base or base in {"not_searched", "not_found", "low_confidence"}:
        return True
    if base == "profile_page":
        return True
    return base not in _PROTECTED_IG_SOURCES


def _load_spa_rows(input_dir: Path, schools: list[str]) -> list[dict[str, str | None]]:
    rows: list[dict[str, str | None]] = []
    for school in schools:
        path = input_dir / SPA_FILES[school]
        if not path.exists():
            log.warning("Missing %s - skipping %s", path, school)
            continue
        payload = json.loads(path.read_text())
        if not isinstance(payload, list):
            raise RuntimeError(f"{path} must contain a JSON list")
        for item in payload:
            rows.append(
                {
                    "School": item.get("School"),
                    "Name": item.get("Name"),
                    "Category": item.get("Category"),
                    "Campus": item.get("Campus"),
                    "Directory URL": item.get("Directory URL"),
                    "Instagram URL": item.get("Instagram URL"),
                    "Instagram Handle": item.get("Instagram Handle"),
                    "IG Source": item.get("IG Source"),
                    "Discord URL": item.get("Discord URL"),
                    "Discord Source": item.get("Discord Source"),
                }
            )
        log.info("Loaded %s rows from %s", len(payload), path.name)
    return rows


def _merge_row(
    existing: dict[str, str | None], spa: dict[str, str | None]
) -> dict[str, str | None]:
    merged = dict(existing)
    if spa.get("Name"):
        merged["Name"] = spa["Name"]
    if spa.get("Category") and not merged.get("Category"):
        merged["Category"] = spa["Category"]
    if spa.get("Campus") and not merged.get("Campus"):
        merged["Campus"] = spa["Campus"]
    if spa.get("Directory URL"):
        merged["Directory URL"] = spa["Directory URL"]

    if _should_update_ig(merged.get("IG Source"), spa.get("IG Source"), spa.get("Instagram URL")):
        merged["Instagram URL"] = spa.get("Instagram URL")
        merged["Instagram Handle"] = spa.get("Instagram Handle")
        merged["IG Source"] = spa.get("IG Source")

    if spa.get("Discord URL") and spa.get("Discord Source") == "profile_page":
        merged["Discord URL"] = spa.get("Discord URL")
        merged["Discord Source"] = spa.get("Discord Source")

    return merged


def merge(input_dir: Path, schools: list[str], xlsx_path: Path) -> tuple[int, int, int]:
    spa_rows = _load_spa_rows(input_dir, schools)
    if not spa_rows:
        raise RuntimeError("No spa rows loaded - run spa_scrape.py first")

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    header = _normalize_header([cell.value for cell in ws[2]])
    if header[: len(COLUMNS)] != COLUMNS:
        if len(header) == 8 and header == COLUMNS[:8]:
            for col_idx, name in enumerate(COLUMNS[8:], start=9):
                ws.cell(row=2, column=col_idx, value=name)
            log.info("Extended xlsx header with Discord columns")
        else:
            raise RuntimeError(f"Unexpected header row: {header!r}")

    existing_by_url: dict[str, int] = {}
    existing_by_school_name: dict[tuple[str, str], int] = {}
    sheet_rows: list[dict[str, str | None]] = []
    for row_idx, raw in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not raw or not any(raw):
            continue
        row = _row_to_dict(raw)
        sheet_rows.append(row)
        url_key = _normalize_url(row.get("Directory URL"))
        if url_key:
            existing_by_url[url_key] = len(sheet_rows) - 1
        school_name_key = ((row.get("School") or "").strip(), (row.get("Name") or "").strip())
        if all(school_name_key):
            existing_by_school_name[school_name_key] = len(sheet_rows) - 1

    updated = 0
    appended = 0
    for spa in spa_rows:
        url_key = _normalize_url(spa.get("Directory URL"))
        if not url_key:
            log.warning("Skipping spa row without Directory URL: %r", spa.get("Name"))
            continue
        school_name_key = (((spa.get("School") or "").strip()), ((spa.get("Name") or "").strip()))
        idx = None
        if url_key and url_key in existing_by_url:
            idx = existing_by_url[url_key]
        elif all(school_name_key) and school_name_key in existing_by_school_name:
            idx = existing_by_school_name[school_name_key]

        if idx is not None:
            before = sheet_rows[idx]
            after = _merge_row(before, spa)
            if after != before:
                sheet_rows[idx] = after
                updated += 1
        else:
            sheet_rows.append(
                {
                    "School": spa.get("School"),
                    "Name": spa.get("Name"),
                    "Category": spa.get("Category"),
                    "Campus": spa.get("Campus"),
                    "Directory URL": spa.get("Directory URL"),
                    "Instagram URL": spa.get("Instagram URL"),
                    "Instagram Handle": spa.get("Instagram Handle"),
                    "IG Source": spa.get("IG Source") or "not_searched",
                    "Discord URL": spa.get("Discord URL"),
                    "Discord Source": spa.get("Discord Source") or "not_searched",
                }
            )
            existing_by_url[url_key] = len(sheet_rows) - 1
            if all(school_name_key):
                existing_by_school_name[school_name_key] = len(sheet_rows) - 1
            appended += 1

    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)
    for row in sheet_rows:
        ws.append(_dict_to_row(row))

    ig_count = sum(1 for row in sheet_rows if row.get("Instagram URL"))
    discord_count = sum(1 for row in sheet_rows if row.get("Discord URL"))
    ws.cell(
        row=1,
        column=1,
        value=(
            f"Master list - {len(sheet_rows):,} student clubs. "
            f"Instagram URLs: {ig_count:,}. Discord URLs: {discord_count:,}. "
            f"SPA merge sources: {', '.join(schools)}."
        ),
    )

    wb.save(xlsx_path)
    return len(sheet_rows), updated, appended


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument(
        "--schools",
        nargs="+",
        choices=sorted(SPA_FILES),
        default=sorted(SPA_FILES),
    )
    parser.add_argument("--xlsx", type=Path, default=XLSX_PATH)
    args = parser.parse_args()

    total, updated, appended = merge(args.input_dir, args.schools, args.xlsx)
    log.info(
        "Wrote %s rows to %s (updated=%s appended=%s)",
        total,
        args.xlsx,
        updated,
        appended,
    )
    return 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    raise SystemExit(main())
