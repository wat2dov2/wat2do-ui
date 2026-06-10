#!/usr/bin/env python3
"""Rewrite Category cells in all_schools_student_clubs_master.xlsx to canonical taxonomy.

Usage (from backend/):
  python scripts/normalize_master_clubs_xlsx.py           # dry-run summary
  python scripts/normalize_master_clubs_xlsx.py --apply     # write workbook
"""

from __future__ import annotations

import argparse
import logging
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.constants.organizations import ORGANIZATION_CATEGORIES
from services.scraper.organization_category_taxonomy import map_directory_category_list

log = logging.getLogger(__name__)

XLSX_PATH = (
    Path(__file__).resolve().parent.parent
    / "services"
    / "scraper"
    / "all_schools_student_clubs_master.xlsx"
)
_CATEGORY_COL = 3  # 1-based column C ("Category")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true", help="Write normalized categories to the workbook.")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    changed = 0
    emptied = 0
    before = Counter()
    after = Counter()
    invalid_after: set[str] = set()

    for row_idx in range(3, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=_CATEGORY_COL)
        old_value = cell.value
        old_text = "" if old_value is None else str(old_value).strip()
        if old_text:
            before[old_text] += 1

        new_categories = map_directory_category_list(old_text if old_text else None)
        for category in new_categories:
            if category not in ORGANIZATION_CATEGORIES:
                invalid_after.add(category)
        new_text = ", ".join(new_categories)
        if new_text:
            after[new_text] += 1
        else:
            emptied += 1

        if new_text != old_text:
            changed += 1
            if args.apply:
                cell.value = new_text or None

    log.info("Rows scanned: %s", ws.max_row - 2)
    log.info("Rows changed: %s", changed)
    log.info("Rows with empty category after normalize: %s", emptied)
    log.info("Unique category values before: %s", len(before))
    log.info("Unique category values after: %s", len(after))

    if invalid_after:
        raise RuntimeError(f"Non-canonical values produced: {sorted(invalid_after)}")

    if args.apply:
        wb.save(XLSX_PATH)
        log.info("Wrote %s", XLSX_PATH)
    else:
        log.info("Dry run only. Pass --apply to write.")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    main()
