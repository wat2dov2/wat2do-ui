#!/usr/bin/env python3
"""Shard the master xlsx for parallel subagent enrichment, then merge results.

Supports Instagram-only search (legacy) and full metadata enrichment (IG,
category, directory URL, Discord, campus).  Each agent writes its own result
JSON; a single ``apply`` pass writes into the xlsx without racing.

Usage:
  # Instagram-only blanks (legacy)
  python backend/scripts/ig_search_shards.py export --shard-size 55
  python backend/scripts/ig_search_shards.py apply

  # Full enrichment: missing IG, category, or directory URL
  python backend/scripts/ig_search_shards.py export --enrich --shard-size 50
  python backend/scripts/ig_search_shards.py apply --enrich

  # Regenerate per-school CSV exports from the master sheet
  python backend/scripts/ig_search_shards.py export-csvs
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.constants.organizations import ORGANIZATION_CATEGORIES

log = logging.getLogger(__name__)

XLSX_PATH = Path(__file__).resolve().parent.parent / "services" / "scraper" / "wat2do-clubs.xlsx"
IG_SHARD_DIR = Path("/tmp/claude/ig_shards")
IG_RESULT_DIR = Path("/tmp/claude/ig_results")
ENRICH_SHARD_DIR = Path("/tmp/claude/enrich_shards")
ENRICH_RESULT_DIR = Path("/tmp/claude/enrich_results")
EXPORT_DIR = Path("/tmp/claude/club_exports")

VALID_CATEGORIES = frozenset(ORGANIZATION_CATEGORIES)

# Global institution accounts that must never be attached to an individual club.
GLOBAL_SCHOOL_HANDLES = frozenset(
    {
        "brocku",
        "carleton_u",
        "carletonu",
        "concordia",
        "concordiauniversity",
        "cornell",
        "cornelluniversity",
        "mcgillu",
        "mcgilluniversity",
        "mcmasteru",
        "mcmasteruniversity",
        "memorialuniversity",
        "munstudents",
        "nyuniversity",
        "nyu",
        "queensuniversity",
        "queensu",
        "simonfraseru",
        "sfu",
        "torontomet",
        "torontomu",
        "ualberta",
        "ubc",
        "ubcnews",
        "ucalgary",
        "umanitoba",
        "uoft",
        "universityoftoronto",
        "uottawa",
        "uofpenn",
        "upenn",
        "westernu",
        "westernuniversity",
        "wlunation",
        "laurier",
        "wilfridlaurieruni",
        "yorku",
        "yorkuniversity",
        "utsc",
        "uoftscarborough",
    }
)

SHARD_FILE_RE = re.compile(r"shard-\d+\.json$")


def _norm(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()


def _canon_ig(url_or_handle: str | None) -> tuple[str | None, str | None]:
    if not url_or_handle:
        return None, None
    text = str(url_or_handle).strip()
    match = re.search(r"instagram\.com/([A-Za-z0-9_.]+)", text)
    handle = match.group(1) if match else text.lstrip("@")
    handle = handle.strip("/").split("?")[0].lower()
    if not handle or handle in {"p", "reel", "explore", "accounts", "stories"}:
        return None, None
    return f"https://www.instagram.com/{handle}/", f"@{handle}"


def _normalize_category(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    matched: list[str] = []
    for category in ORGANIZATION_CATEGORIES:
        if category in text and category not in matched:
            matched.append(category)
    if not matched:
        return None
    return ", ".join(matched)


def _normalize_url(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text.startswith(("http://", "https://")):
        return None
    return text


def _normalize_str(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _school_slug(school: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", school.lower().replace("'", "")).strip("_")


def _load_missing_ig(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=3, values_only=True):
        if not raw or not raw[0] or not raw[1]:
            continue
        if raw[5]:
            continue
        rows.append(
            {
                "school": str(raw[0]).strip(),
                "name": str(raw[1]).strip(),
                "directory_url": str(raw[4]).strip() if raw[4] else None,
            }
        )
    wb.close()
    return rows


def _load_enrichment_rows(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=3, values_only=True):
        if not raw or not raw[0] or not raw[1]:
            continue
        needs_ig = not raw[5]
        needs_category = not raw[2]
        needs_directory = not raw[4]
        needs_discord = not raw[8]
        if not (needs_ig or needs_category or needs_directory):
            continue
        rows.append(
            {
                "school": str(raw[0]).strip(),
                "name": str(raw[1]).strip(),
                "category": str(raw[2]).strip() if raw[2] else None,
                "campus": str(raw[3]).strip() if raw[3] else None,
                "directory_url": str(raw[4]).strip() if raw[4] else None,
                "instagram_url": str(raw[5]).strip() if raw[5] else None,
                "discord_url": str(raw[8]).strip() if raw[8] else None,
                "needs": {
                    "instagram": needs_ig,
                    "category": needs_category,
                    "directory_url": needs_directory,
                    "discord": needs_discord,
                },
            }
        )
    wb.close()
    return rows


def export_shards(
    xlsx_path: Path,
    shard_size: int,
    *,
    enrich: bool,
) -> None:
    rows = _load_enrichment_rows(xlsx_path) if enrich else _load_missing_ig(xlsx_path)
    shard_dir = ENRICH_SHARD_DIR if enrich else IG_SHARD_DIR
    result_dir = ENRICH_RESULT_DIR if enrich else IG_RESULT_DIR
    shard_dir.mkdir(parents=True, exist_ok=True)
    result_dir.mkdir(parents=True, exist_ok=True)
    for existing in shard_dir.glob("shard-*.json"):
        existing.unlink()
    for existing in result_dir.glob("shard-*.json"):
        existing.unlink()

    shards = [rows[i : i + shard_size] for i in range(0, len(rows), shard_size)]
    width = max(2, len(str(len(shards) - 1)))
    for idx, shard in enumerate(shards):
        path = shard_dir / f"shard-{idx:0{width}d}.json"
        path.write_text(json.dumps(shard, indent=2, ensure_ascii=False) + "\n")
    log.info(
        "Exported %s %s rows into %s shards (size=%s) under %s",
        len(rows),
        "enrichment" if enrich else "missing-IG",
        len(shards),
        shard_size,
        shard_dir,
    )


def _collect_results(result_dir: Path) -> tuple[dict[tuple[str, str], dict], dict[str, int]]:
    result_files = sorted(p for p in result_dir.glob("shard-*.json") if SHARD_FILE_RE.match(p.name))
    if not result_files:
        raise SystemExit(f"No result files under {result_dir}")

    matches: dict[tuple[str, str], dict] = {}
    stats = {"seen": 0, "skipped_global": 0, "skipped_bad_ig": 0}
    for path in result_files:
        payload = json.loads(path.read_text())
        items = (
            payload
            if isinstance(payload, list)
            else payload.get("matches", payload.get("results", []))
        )
        for item in items:
            stats["seen"] += 1
            key = (_norm(item.get("school")), _norm(item.get("name")))
            ig_url = ig_handle = None
            if item.get("instagram_url") or item.get("instagram_handle"):
                ig_url, ig_handle = _canon_ig(
                    item.get("instagram_url") or item.get("instagram_handle")
                )
                if not ig_url or not ig_handle:
                    stats["skipped_bad_ig"] += 1
                elif ig_handle.lstrip("@") in GLOBAL_SCHOOL_HANDLES:
                    stats["skipped_global"] += 1
                    ig_url = ig_handle = None

            category = _normalize_category(item.get("category"))
            directory_url = _normalize_url(item.get("directory_url"))
            discord_url = _normalize_url(item.get("discord_url"))
            campus = _normalize_str(item.get("campus"))
            source = item.get("source") or "web_search_enrichment_pass"

            if not any([ig_url, category, directory_url, discord_url, campus]):
                continue

            matches[key] = {
                "instagram_url": ig_url,
                "instagram_handle": ig_handle,
                "category": category,
                "directory_url": directory_url,
                "discord_url": discord_url,
                "campus": campus,
                "source": source,
            }
    return matches, stats


def apply(xlsx_path: Path, *, enrich: bool) -> None:
    result_dir = ENRICH_RESULT_DIR if enrich else IG_RESULT_DIR
    matches, stats = _collect_results(result_dir)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    owner_of_handle: dict[str, tuple[str, str]] = {}
    for row in range(3, ws.max_row + 1):
        existing_url = ws.cell(row=row, column=6).value
        if not existing_url:
            continue
        _, existing_handle = _canon_ig(existing_url)
        if existing_handle:
            owner_of_handle.setdefault(
                existing_handle.lstrip("@"),
                (_norm(ws.cell(row=row, column=1).value), _norm(ws.cell(row=row, column=2).value)),
            )

    applied_ig = applied_category = applied_directory = applied_discord = applied_campus = 0
    skipped_dupe = 0
    for row in range(3, ws.max_row + 1):
        key = (_norm(ws.cell(row=row, column=1).value), _norm(ws.cell(row=row, column=2).value))
        match = matches.get(key)
        if not match:
            continue

        if not ws.cell(row=row, column=6).value and match.get("instagram_url"):
            handle_key = match["instagram_handle"].lstrip("@")
            owner = owner_of_handle.get(handle_key)
            if owner and owner != key:
                skipped_dupe += 1
            else:
                ws.cell(row=row, column=6, value=match["instagram_url"])
                ws.cell(row=row, column=7, value=match["instagram_handle"])
                ws.cell(row=row, column=8, value=f"confirmed|{match['source']}")
                owner_of_handle[handle_key] = key
                applied_ig += 1

        if not ws.cell(row=row, column=3).value and match.get("category"):
            ws.cell(row=row, column=3, value=match["category"])
            applied_category += 1

        if not ws.cell(row=row, column=4).value and match.get("campus"):
            ws.cell(row=row, column=4, value=match["campus"])
            applied_campus += 1

        if not ws.cell(row=row, column=5).value and match.get("directory_url"):
            ws.cell(row=row, column=5, value=match["directory_url"])
            applied_directory += 1

        if not ws.cell(row=row, column=9).value and match.get("discord_url"):
            ws.cell(row=row, column=9, value=match["discord_url"])
            ws.cell(row=row, column=10, value=f"confirmed|{match['source']}")
            applied_discord += 1

    row_count = ws.max_row - 2
    ig_count = sum(1 for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=6).value)
    cat_count = sum(1 for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=3).value)
    discord_count = sum(1 for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=9).value)
    ws.cell(
        row=1,
        column=1,
        value=(
            f"Master list - {row_count:,} student clubs. "
            f"Instagram URLs: {ig_count:,}. Categories: {cat_count:,}. "
            f"Discord URLs: {discord_count:,}. "
            f"SPA merge sources: {'enrichment' if enrich else 'web_search'}."
        ),
    )
    wb.save(xlsx_path)
    log.info(
        "Applied IG=%s category=%s directory=%s discord=%s campus=%s "
        "(candidates=%s, skipped_global=%s, skipped_bad_ig=%s, skipped_dupe=%s). "
        "Sheet now has %s IG / %s categories / %s rows.",
        applied_ig,
        applied_category,
        applied_directory,
        applied_discord,
        applied_campus,
        stats["seen"],
        stats["skipped_global"],
        stats["skipped_bad_ig"],
        skipped_dupe,
        ig_count,
        cat_count,
        row_count,
    )


def export_csvs(xlsx_path: Path) -> None:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    by_school: dict[str, list[dict]] = {}
    fieldnames = [
        "name",
        "category",
        "campus",
        "directory_url",
        "instagram_handle",
        "instagram_url",
        "ig_source",
        "discord_url",
        "discord_source",
    ]
    for raw in ws.iter_rows(min_row=3, values_only=True):
        if not raw or not raw[0] or not raw[1]:
            continue
        school = str(raw[0]).strip()
        handle = (raw[6] or "").strip()
        if handle and not handle.startswith("@"):
            handle = "@" + handle
        if not handle and raw[5]:
            handle = str(raw[5]).strip()
        by_school.setdefault(school, []).append(
            {
                "name": raw[1] or "",
                "category": raw[2] or "",
                "campus": raw[3] or "",
                "directory_url": raw[4] or "",
                "instagram_handle": handle,
                "instagram_url": raw[5] or "",
                "ig_source": raw[7] or "",
                "discord_url": raw[8] or "",
                "discord_source": raw[9] or "",
            }
        )
    wb.close()

    for school, rows in by_school.items():
        rows.sort(key=lambda r: r["name"].lower())
        out = EXPORT_DIR / f"{_school_slug(school)}_clubs_instagram.csv"
        with out.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        log.info("Wrote %s (%s clubs)", out.name, len(rows))
    log.info("Exported %s school CSV files to %s", len(by_school), EXPORT_DIR)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=XLSX_PATH)
    sub = parser.add_subparsers(dest="command", required=True)

    export_parser = sub.add_parser("export", help="Split rows into shard files")
    export_parser.add_argument("--shard-size", type=int, default=50)
    export_parser.add_argument(
        "--enrich",
        action="store_true",
        help="Export rows missing IG, category, or directory URL (full metadata pass)",
    )

    apply_parser = sub.add_parser("apply", help="Merge subagent result files into the xlsx")
    apply_parser.add_argument(
        "--enrich",
        action="store_true",
        help="Read from enrich_results/ and apply all metadata fields",
    )

    sub.add_parser("export-csvs", help="Regenerate per-school CSV exports from the master sheet")

    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    if args.command == "export":
        export_shards(args.xlsx, args.shard_size, enrich=args.enrich)
    elif args.command == "apply":
        apply(args.xlsx, enrich=args.enrich)
    elif args.command == "export-csvs":
        export_csvs(args.xlsx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
